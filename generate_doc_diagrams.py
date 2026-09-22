import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import base64
import zlib
import urllib.request
import os

def download_diagram(diagram_type, diagram_text, filename):
    compressed = zlib.compress(diagram_text.encode('utf-8'), 9)
    encoded = base64.urlsafe_b64encode(compressed).decode('utf-8')
    url = f"https://kroki.io/{diagram_type}/png/{encoded}"
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    try:
        with urllib.request.urlopen(req) as response, open(filename, 'wb') as out_file:
            out_file.write(response.read())
    except Exception as e:
        print(f"Failed to download {diagram_type}: {e}")
        # Create a dummy image if download fails to not crash the script
        with open(filename, 'wb') as f:
            pass # We'll check size before adding

plantuml_1 = """@startuml
skinparam componentStyle rectangle
node "Guest Browser" as Client
package "BFF Layer (Next.js)" {
  [app/api/app/*] as BFF_API
}
package "API Gateway (Fastify)" {
  [/api/v1/*] as Gateway_API
}
database "Firestore DB" as DB

Client --> BFF_API : HTTP Request
BFF_API --> Gateway_API : Zod Validation & DTO
Gateway_API --> DB : Compare-and-set
DB --> Gateway_API
Gateway_API --> BFF_API : Flat Envelope
BFF_API --> Client : {ok, data, error, meta}
@enduml"""

plantuml_2 = """@startuml
actor Client
participant Gateway
database "Firestore DB" as DB

Client -> Gateway : POST /api/v2/checkout/initiate
Gateway -> Gateway : rateLimit -> validateV2 -> authorize
Gateway -> DB : runTransaction()
note right of DB
  Compare-and-Set:
  Write Order + Outbox Event
end note
DB --> Gateway : Commit Success
Gateway --> Client : 200 OK (Flat Envelope)
@enduml"""

print("Downloading diagrams...")
download_diagram('plantuml', plantuml_1, 'arch_diag.png')
download_diagram('plantuml', plantuml_2, 'seq_diag.png')

print("Building docx...")
doc = docx.Document()
title = doc.add_heading('System Design: Guest Portal BFF Architecture and Incremental Migration Strategy', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('1. Executive Summary & Context', level=1)
doc.add_paragraph("This document outlines the architectural transition of the Guest Portal from a monolithic direct-to-database pattern to a modernized Backend-for-Frontend (BFF) approach. Historically, the Guest Portal relied on thick client-side logic to aggregate data directly from our primary datastores. To resolve this, we are introducing a thin UI layer supported by localized BFF read-models.")

doc.add_heading('2. Architectural Goals & Constraints', level=1)
p = doc.add_paragraph(style='List Bullet')
p.add_run("Strict Boundary Enforcement: ").bold = True
p.add_run("The Guest Portal's BFF handlers must not bypass the core API gateway. The gateway remains the single source of truth.")
p = doc.add_paragraph(style='List Bullet')
p.add_run("Standardized Flat Envelopes: ").bold = True
p.add_run("Every new BFF endpoint must strictly adhere to our unified response envelope: { ok, data, error, meta }. All error envelopes are flat across the entire stack.")
p = doc.add_paragraph(style='List Bullet')
p.add_run("Compare-and-Set Durability: ").bold = True
p.add_run("All state mutations implement compare-and-set transactions (e.g., verifying expectedVersion) to prevent lost updates during concurrent checkout flows.")

doc.add_heading('3. System Design: The BFF Pattern', level=1)
doc.add_paragraph("The localized BFF layer is built directly into the Next.js application using standard API route handlers under app/api/app/*. Behind the scenes, the BFF handler validates the incoming request DTO using Zod. Once validated, it orchestrates calls to the backend API Gateway.")

if os.path.getsize('arch_diag.png') > 0:
    doc.add_picture('arch_diag.png', width=Inches(5.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Figure 1: High-level request flow through the BFF and API Gateway layers.", style='Caption')

doc.add_heading('4. Data Flow & Transactional Integrity', level=1)
doc.add_paragraph("For complex flows such as guest checkout and ticket reservation, we implement a Transactional Outbox pattern alongside Compare-and-Set lock TTLs. Every order creation runs strictly inside a Firestore runTransaction block, persisting the business write and outbox event atomically.")

if os.path.getsize('seq_diag.png') > 0:
    doc.add_picture('seq_diag.png', width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Figure 2: Transactional Checkout Flow.", style='Caption')

doc.add_heading('5. Edge Cases, Failure Modes, and Mitigation', level=1)
p = doc.add_paragraph(style='List Bullet')
p.add_run("Policy Ordering: ").bold = True
p.add_run("Middleware strictly executes in order: rate-limit -> validate -> authorize -> cache. This ensures malformed DTOs are rejected with 422 before ABAC authorization logic wastes DB reads.")
p = doc.add_paragraph(style='List Bullet')
p.add_run("Gateway Latency Spikes: ").bold = True
p.add_run("If the upstream API Gateway experiences degraded performance, the BFF layer enforces strict 5-second timeout policies for read operations.")

doc.add_heading('6. Migration Strategy & Rollout Plan', level=1)
doc.add_paragraph("During Phase 1, the new BFF endpoints will be deployed to production in 'shadow mode'. Once parity is achieved and error rates drop below our 0.1% threshold, Phase 2 will involve a gradual traffic shift, starting with 5% of users.")

doc.save('Guest_Portal_BFF_Architecture.docx')
print("Done!")
